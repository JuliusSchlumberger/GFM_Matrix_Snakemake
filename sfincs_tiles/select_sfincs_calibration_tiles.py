"""Selects tiles for the "sfincs_calibration" study (2026-10, supersedes the
validation_sfincs_v1-v6 studies' stricter eligibility-based selection,
select_validation_tiles.py): hop_distance==0 tiles, globally stratified -
no other eligibility/representativeness criteria applied (explicit user
direction - drops the old script's max-cells/river-fraction/COAST-HG-
station-match eligibility filter, min-ocean-frac floor, high-latitude
cutoff, and model_outputs-already-exists requirement entirely).

Antimeridian-crossing tiles are still excluded - this is NOT a
representativeness criterion, it is a hard technical necessity:
hydromt_sfincs's own water_level.create() masking does a shapely
union_all() in raw EPSG:4326 lon/lat and crashes with a TopologyException
on a dateline-wrapping geometry, regardless of buffer tuning (see
build_sfincs_tile.py's _classify_water_level_create_error). There is no
equivalent hard constraint for the dropped criteria above.

Default tile count is round(25% x the FULL production tile grid), e.g.
1,126 of the grid's 4,504 tiles as of 2026-10 (not 25% of the hop=0 subset
alone) - all drawn from the grid's 3,761 hop_distance==0 tiles, comfortably
inside that pool.

Reuses select_validation_tiles.py's plotting/sampling helpers directly
(is_antimeridian_tile, stratified_sample_exact, _area_km2, read_ocean_frac,
plot_tile_locations_map, plot_selection_histograms) - only the selection
logic itself (which filters get applied, and the stats-table labels, which
would otherwise misleadingly describe a now-nonexistent eligibility stage)
differs from that script.

Writes tile_ids.txt (selected tile IDs) under {root}/{base_dir_name}/, plus
tile_selection_metadata.csv, tile_locations_map.png,
tile_selection_histograms.png, tile_selection_statistics.xlsx (ocean_frac/
area are informational columns on these outputs only, never a filter).

Usage:
    python select_sfincs_calibration_tiles.py
    python select_sfincs_calibration_tiles.py --n-tiles 1126
"""

from __future__ import annotations

import sys
from pathlib import Path

import geopandas as gpd
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from select_validation_tiles import (  # noqa: E402
    _area_km2,
    is_antimeridian_tile,
    plot_selection_histograms,
    plot_tile_locations_map,
    read_ocean_frac,
    stratified_sample_exact,
)
from select_test_tiles import GRID_DEG  # noqa: E402
from gfm_config import read_root  # noqa: E402

ANTIMERIDIAN_BBOX_WIDTH_DEG_DEFAULT = 60.0  # wider than any real tile's bbox
BASE_DIR_NAME_DEFAULT = "sfincs_calibration"
TILE_FRACTION_DEFAULT = 0.25


def write_statistics_excel(
    tiles_gdf: gpd.GeoDataFrame, hop0: gpd.GeoDataFrame, selected: pd.DataFrame, out_path: Path,
) -> None:
    """Writes tile-count/size/ocean-fraction summary tables - own version
    (not select_validation_tiles.py's), since that one's labels describe an
    eligibility stage ("Eligible pool (size/river/COAST-HG criteria met)")
    that doesn't exist in this simplified selection.
    """
    counts_df = pd.DataFrame([
        {"Population": "Full production grid", "n_tiles": len(tiles_gdf)},
        {"Population": "hop_distance==0, antimeridian-excluded", "n_tiles": len(hop0)},
        {"Population": "Selected for sfincs_calibration", "n_tiles": len(selected)},
    ])

    def _row(label: str, df: pd.DataFrame) -> dict:
        area = df["area_km2_bbox"].to_numpy()
        ocean = df["ocean_frac"].dropna().to_numpy()
        row = {
            "Population": label, "n_tiles": len(df),
            "area_km2_min": float(area.min()), "area_km2_median": float(pd.Series(area).median()),
            "area_km2_mean": float(area.mean()), "area_km2_max": float(area.max()),
        }
        if len(ocean) > 0:
            row.update({
                "ocean_frac_min": float(ocean.min()), "ocean_frac_median": float(pd.Series(ocean).median()),
                "ocean_frac_mean": float(ocean.mean()), "ocean_frac_max": float(ocean.max()),
            })
        return row

    stats_df = pd.DataFrame([
        _row("hop_distance==0, antimeridian-excluded", hop0),
        _row("Selected", selected),
    ])

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        counts_df.to_excel(writer, sheet_name="Tile counts", index=False)
        stats_df.to_excel(writer, sheet_name="Size and ocean fraction", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(out_path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col_cells in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells) + 2
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(width, 10), 40)
    wb.save(out_path)


def main() -> None:
    import argparse

    _repo_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", default=str(_repo_root / "snakemake_workflow" / "config" / "config.yml"))
    parser.add_argument(
        "--n-tiles", type=int, default=None,
        help=f"defaults to round({TILE_FRACTION_DEFAULT:.0%} * full production grid size)",
    )
    parser.add_argument("--antimeridian-width-deg", type=float, default=ANTIMERIDIAN_BBOX_WIDTH_DEG_DEFAULT)
    parser.add_argument("--base-dir-name", default=BASE_DIR_NAME_DEFAULT)
    parser.add_argument("--seed", type=int, default=0)
    args = parser.parse_args()

    root = read_root(Path(args.config))
    tiles_gdf = gpd.read_file(root / "processed_inputs" / "mask" / "domain_tiles_global.gpkg")
    n_all = len(tiles_gdf)
    n_tiles = args.n_tiles if args.n_tiles is not None else round(TILE_FRACTION_DEFAULT * n_all)

    hop0 = tiles_gdf[tiles_gdf["hop_distance"] == 0].copy()
    print(f"{n_all} tiles total, {len(hop0)} with hop_distance == 0; target {n_tiles} tile(s) "
          f"({'explicit' if args.n_tiles is not None else f'{TILE_FRACTION_DEFAULT:.0%} of {n_all}'})")

    hop0["antimeridian"] = hop0.geometry.apply(lambda g: is_antimeridian_tile(g, args.antimeridian_width_deg))
    n_antimeridian = int(hop0["antimeridian"].sum())
    hop0 = hop0[~hop0["antimeridian"]].copy()
    print(f"{n_antimeridian} antimeridian-crossing tile(s) excluded (hard technical necessity, not a "
          f"representativeness criterion - see module docstring), {len(hop0)} remain")

    if n_tiles > len(hop0):
        raise ValueError(f"requested {n_tiles} tile(s) but only {len(hop0)} hop_distance==0 tile(s) available")

    centroid = hop0.geometry.centroid
    hop0["lon"] = centroid.x.to_numpy()
    hop0["lat"] = centroid.y.to_numpy()
    hop0["area_km2_bbox"] = _area_km2(hop0)

    print(f"Reading ocean_frac for all {len(hop0)} tile(s) (informational only, not a selection filter)...", flush=True)
    hop0["ocean_frac"] = hop0["tile_id"].apply(lambda t: read_ocean_frac(int(t), root))

    selected = stratified_sample_exact(hop0, n_tiles, GRID_DEG, args.seed)
    print(f"\nSelected: {len(selected)} tiles (hop_distance==0, antimeridian-excluded, globally "
          f"stratified - no other criteria)")

    out_dir = root / args.base_dir_name
    out_dir.mkdir(parents=True, exist_ok=True)
    selected.to_csv(out_dir / "tile_selection_metadata.csv", index=False)
    (out_dir / "tile_ids.txt").write_text("\n".join(str(t) for t in selected["tile_id"]) + "\n")
    print(f"\nWrote {out_dir / 'tile_selection_metadata.csv'}, {out_dir / 'tile_ids.txt'} ({len(selected)})")

    map_path = out_dir / "tile_locations_map.png"
    plot_tile_locations_map(selected, map_path)
    print(f"Wrote {map_path}")

    hist_path = out_dir / "tile_selection_histograms.png"
    plot_selection_histograms(selected, hop0, hist_path)
    print(f"Wrote {hist_path}")

    xlsx_path = out_dir / "tile_selection_statistics.xlsx"
    write_statistics_excel(tiles_gdf, hop0, selected, xlsx_path)
    print(f"Wrote {xlsx_path}")


if __name__ == "__main__":
    main()
